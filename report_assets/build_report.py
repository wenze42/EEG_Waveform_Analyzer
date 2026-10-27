# -*- coding: utf-8 -*-
"""生成《负反馈放大器》实验报告 docx"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/workspace/负反馈放大器实验报告.docx"
ASSETS = "/workspace/report_assets"

doc = Document()

# ---------- 页面设置 ----------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.3)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.4)

# ---------- 默认样式 ----------
normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)


def set_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def body(text, indent=True, space_after=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_together = True
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    r = p.add_run(text)
    set_font(r, "宋体", 12)
    return p


def formula(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, "宋体", 11)
    return p


def heading(text, level=1):
    h = doc.add_heading(level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    set_font(r, "黑体", 15 if level == 1 else 12.5, True, RGBColor(0, 0, 0))


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, "宋体", 10.5)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(headers, rows, widths, font_size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(h)
        set_font(r, "宋体", font_size, True)
        shade(c, "F2F2F2")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.width = Cm(widths[i])
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            set_font(r, "宋体", font_size)
    # 表头跨页重复
    trPr0 = t.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr0.append(tblHeader)
    # 行高
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "360")
        trPr.append(trHeight)
    return t


def footer_page_number():
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    set_font(r, "宋体", 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r2 = OxmlElement("w:r")
    t2 = OxmlElement("w:t")
    t2.text = "1"
    r2.append(t2)
    fld.append(r2)
    p._p.append(fld)
    r3 = p.add_run(" 页")
    set_font(r3, "宋体", 9)


# =====================================================================
# 封面
# =====================================================================
for _ in range(2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("南开大学电子信息与光学工程学院")
set_font(r, "宋体", 20, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
r = p.add_run("实 验 报 告")
set_font(r, "宋体", 36, True)

for _ in range(4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

def field_para(label, value, tail=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(2.4)
    p.paragraph_format.space_before = Pt(11)
    r = p.add_run(label + "　")
    set_font(r, "宋体", 12.5)
    gap = p.add_run("　")
    set_font(gap, "宋体", 12.5)
    r2 = p.add_run(value + "　")
    set_font(r2, "宋体", 12.5)
    r2.underline = True
    if tail:
        r3 = p.add_run(" " + tail)
        set_font(r3, "宋体", 12.5)

field_para("实验序号", "3")
field_para("实验名称", "负反馈放大器")

p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(2.4)
p.paragraph_format.space_before = Pt(11)
r = p.add_run("　")
set_font(r, "宋体", 12.5)
r = p.add_run("电子信息与光学工程学院")
set_font(r, "宋体", 12.5)
r.underline = True
r = p.add_run("学院　")
set_font(r, "宋体", 12.5)
r = p.add_run("________")
set_font(r, "宋体", 12.5)
r = p.add_run(" 系")
set_font(r, "宋体", 12.5)

p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(2.4)
p.paragraph_format.space_before = Pt(11)
for label in ["姓名", "学号", "实验台号"]:
    r = p.add_run(label + "　")
    set_font(r, "宋体", 12.5)
    r = p.add_run("________")
    set_font(r, "宋体", 12.5)
    r = p.add_run("　")
    set_font(r, "宋体", 12.5)

field_para("实验日期及时间", "2026年10月27日　星期二　第3～4节")
field_para("实验地点", "模拟电路实验室")

doc.add_page_break()

# =====================================================================
# 一、实验目的
# =====================================================================
heading("一、实验目的")
purposes = [
    "1. 掌握两级阻容耦合负反馈放大电路静态工作点的测量与调试方法，能根据测量结果判断晶体管的工作状态。",
    "2. 学习负反馈放大器交流性能指标（电压放大倍数、反馈系数、输入电阻、输出电阻、通频带）的测量方法。",
    "3. 通过开环与闭环实验结果的对比，理解电压串联负反馈对放大器增益稳定性、输入电阻、输出电阻及通频带的影响。",
    "4. 验证深度负反馈条件下 Avf≈1/Fv 的近似关系，加深对负反馈改善放大器性能的理解。",
]
for s in purposes:
    body(s, indent=False, space_after=2)

# =====================================================================
# 二、实验仪器及元器件
# =====================================================================
heading("二、实验仪器及元器件")
caption("表2-1  实验仪器")
table(
    ["实验仪器", "型号", "备注"],
    [
        ["双踪数字示波器", "TDS 1012C", "1台，AC耦合测量交流量"],
        ["函数信号发生器", "ATF20B", "1台，输出正弦信号"],
        ["数字万用表", "VC890D", "1台，DCV档测静态工作点"],
        ["双路直流稳压电源", "GPS-3030D", "1台，输出+12V"],
        ["模拟电路实验箱", "—", "1台"],
    ],
    [4.4, 3.8, 7.6],
)

caption("表2-2  电子元器件")
table(
    ["电子元器件", "规格", "数量"],
    [
        ["晶体管", "9013", "2只"],
        ["电阻", "1MΩ", "1只"],
        ["电阻", "200kΩ", "1只"],
        ["电阻", "2kΩ", "2只"],
        ["电阻", "30kΩ", "1只"],
        ["电阻", "10kΩ", "1只"],
        ["电阻", "1kΩ", "1只"],
        ["电阻", "200Ω", "1只"],
        ["电阻", "5.1kΩ", "2只"],
        ["电解电容", "10μF", "3只"],
        ["电解电容", "100μF", "1只"],
    ],
    [4.6, 4.0, 4.0],
)

# =====================================================================
# 三、实验原理
# =====================================================================
heading("三、实验原理")

body("实验电路如图3-1所示，为两级阻容耦合共射放大电路。T1、T2均采用9013型NPN晶体管。第一级由RB1提供基极偏置，RC1为集电极负载，RE1为发射极电阻；第二级由RB12、RB22分压提供基极偏置，RC2为集电极负载，RE2为发射极电阻，CE2为发射极旁路电容。级间通过C2耦合，输出经C3耦合后接负载RL。")

body("反馈电阻RF从输出端引回T1的发射极，与RE1共同组成反馈网络。反馈信号以电压形式与输入信号在T1发射极串联相减，因此该电路为电压串联负反馈。由于反馈电压取自输出电压，反馈系数近似为")
formula("Fv = Vf / Vo ≈ RE1 / (RE1 + RF) = 200 / (200 + 5100) ≈ 0.0377")

body("闭环电压放大倍数为")
formula("Avf = Vo / Vi = A / (1 + AF)")
body("式中A为开环电压放大倍数，F为反馈系数，(1+AF)为反馈深度。当满足深度负反馈条件 AF>>1 时，Avf≈1/Fv，闭环增益主要取决于反馈网络，而与晶体管参数关系不大，因而增益稳定性显著提高。")

body("负反馈对放大器输入电阻、输出电阻和通频带的影响可表示为：")
formula("Rif = Ri(1 + AF)，Rof = Ro / (1 + AF)")
formula("fHf = (1 + AF) fH，fLf = fL / (1 + AF)，fBf = fHf - fLf")
body("即电压串联负反馈使输入电阻增大、输出电阻减小、通频带展宽，而增益带宽积近似保持不变。")

body("静态工作点的计算关系为：")
formula("IB1 = (VCC - VB1) / RB1，IC1 = (VCC - VC1) / RC1，VBE1 = VB1 - VE1，VCE1 = VC1 - VE1")
formula("IE2 = VE2 / RE2，IC2 ≈ (VCC - VC2) / RC2，IB2 = (VCC - VB2)/RB12 - VB2/RB22")
formula("VBE2 = VB2 - VE2，VCE2 = VC2 - VE2")
body("晶体管工作在放大区的条件是发射结正偏（VBE≈0.6～0.7 V）、集电结反偏（VCE>0.3 V）。")

body("交流参数的测量方法如下：")
body("（1）输入电阻采用串联电阻法。在信号源与放大器输入端之间串入已知电阻R，分别测量R两端对地的电压Vs和Vi，则")
formula("Ri = Vi / (Vs - Vi) × R")
body("（2）输出电阻采用负载变化法。分别测量不接负载时的输出电压Vo∞和接入负载RL后的输出电压VoL，则")
formula("Ro = (Vo∞ / VoL - 1) × RL")
body("（3）通频带采用0.707法。在中断频率附近测得中频输出电压Vom，保持输入电压不变，改变频率使输出电压下降到0.707Vom，对应上限频率fH和下限频率fL，通频带 fB=fH-fL。")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.add_run().add_picture(os.path.join(ASSETS, "circuit.png"), width=Cm(15.2))
caption("图3-1  负反馈放大器实验电路")

# =====================================================================
# 四、实验内容、步骤与要求
# =====================================================================
heading("四、实验内容、步骤与要求")
body("1. 静态工作点的测量与调试", indent=False)
body("按图3-1连接电路，检查电源极性和连线无误后接通+12 V电源。将输入端对地短路，保证输入交流信号为零。用万用表DCV档分别测量B1、E1、C1和B2、E2、C2对地的直流电压，记录并计算各静态参数。参考值为：VB1≈1V、VC1≈8.4V、VB2≈3V、VC2≈7.6V。若偏差过大，应检查12 V电源、电阻阻值、电容极性及连线是否正确；静态工作点正常后方可进行后续实验。")

body("2. 闭环电压放大倍数和反馈系数的测量", indent=False)
body("接入RL=5.1 kΩ，调节信号发生器输出1 kHz正弦波，使输入电压Vi约为50 mV（峰峰值，下同）。用示波器AC耦合方式分别测量Vi、Vo和Vf（Vf为RE1两端电压），计算Avf=Vo/Vi、Fv=Vf/Vo，并与1/Fv比较。")

body("3. 闭环输入电阻的测量", indent=False)
body("去掉RL，在信号源与放大器输入端之间串联R=200 kΩ。调节信号源使Vi约为100 mV（1 kHz），用示波器测量Vs和Vi，按Ri=Vi/(Vs-Vi)×R计算闭环输入电阻。")

body("4. 闭环输出电阻的测量", indent=False)
body("接入RL=200 Ω，输入Vi=50 mV、f=1 kHz。分别测量接入负载前、后的输出电压Vo∞和VoL，按Ro=(Vo∞/VoL-1)×RL计算闭环输出电阻。")

body("5. 闭环通频带的测量", indent=False)
body("接入RL=5.1 kΩ，Vi=50 mV、f=1 kHz，测量中频输出电压Vom。保持Vi不变，缓慢升高频率，当Vo下降到0.707Vom时记录上限频率fH；再降低频率，当Vo下降到0.707Vom时记录下限频率fL。计算fB=fH-fL。")

body("6. 开环参数的测量", indent=False)
body("断开反馈电阻RF，使放大器处于开环状态。由于开环增益很高，输入信号应减小到10 mV左右，以免输出波形失真。保持RL=5.1 kΩ，测量开环电压放大倍数；串联R=5.1 kΩ、取Vi≈100 mV测量开环输入电阻；保持RL=5.1 kΩ测量开环输出电阻；按0.707法测量开环上限频率和下限频率。每次改变频率时均应保持Vi不变。")

body("7. 注意事项", indent=False)
notes = [
    "（1）测量静态工作点时，必须保证输入端交流信号为零，否则直流读数会偏离真实静态值。",
    "（2）除静态工作点外，Vi、Vo、Vf、Vs均用示波器AC耦合方式测量峰峰值。",
    "（3）Vi的峰峰值一般小于信号源显示的设定值，应以示波器实测值为准。",
    "（4）测量输入电阻时应去掉RL；闭环时串联200 kΩ，开环时串联5.1 kΩ。",
    "（5）测量Vo、Vf时，示波器探头地线必须接在电路的GND端。",
    "（6）频率过高或过低时波形可能不再是标准正弦波，此时只记录与Vi同频的峰峰值；应适当调整示波器的时基和垂直档位，使波形显示几个完整周期并尽量充满屏幕。",
    "（7）连线应疏密适度，尽量少用长导线，以减小寄生耦合和干扰。",
]
for s in notes:
    body(s, indent=False, space_after=1)

# =====================================================================
# 五、实验结果
# =====================================================================
heading("五、实验结果")

heading("5.1  静态工作点", 2)
body("在输入交流信号为零的条件下，用数字万用表DCV档测得各电极对地电压如表5-1所示。")
caption("表5-1  静态工作点电压测量值")
table(
    ["参数", "VB1", "VE1", "VC1", "VB2", "VE2", "VC2"],
    [["测量值/V", "1.173", "0.522", "6.747", "2.928", "2.282", "7.665"]],
    [2.0, 2.1, 2.1, 2.1, 2.1, 2.1, 2.1],
)
body("由测量值计算的静态参数列于表5-2。")
caption("表5-2  静态工作点计算值")
table(
    ["晶体管", "VBEQ/V", "VCEQ/V", "ICQ/mA", "IBQ/μA", "β"],
    [
        ["T1", "0.651", "6.225", "2.63", "10.8", "243"],
        ["T2", "0.646", "5.383", "2.17", "9.6", "226"],
    ],
    [2.3, 2.9, 2.9, 2.8, 2.8, 2.1],
)
body("计算过程如下：", indent=False)
formula("T1：IB1Q=(12-1.173)/1 MΩ≈10.8 μA，IC1Q=(12-6.747)/2 kΩ≈2.63 mA")
formula("VBE1Q=1.173-0.522=0.651 V，VCE1Q=6.747-0.522=6.225 V")
formula("T2：IE2Q=2.282/1 kΩ=2.28 mA，IC2Q=(12-7.665)/2 kΩ≈2.17 mA")
formula("IB2Q=(12-2.928)/30 kΩ-2.928/10 kΩ≈9.6 μA，VBE2Q=0.646 V，VCE2Q=5.383 V")
body("由表5-2可见，两只晶体管的VBE均约为0.65 V，发射结正偏；VCE均大于5 V，集电结反偏，说明T1、T2均工作在放大区。VB2、VC2与参考值（约3 V、7.6 V）吻合较好；VC1实测为6.747 V，低于参考值8.4 V，主要原因是所用9013晶体管的β偏大（约243）以及RB1、RC1存在容差，使第一级集电极电流偏大。电路输出波形无明显失真，说明静态工作点仍可正常工作。")

doc.add_page_break()
heading("5.2  闭环特性", 2)
body("接入反馈电阻RF后，在RL=5.1 kΩ、f=1 kHz条件下测得Vi=46.7 mV、Vo=1.175 V、Vf=36.4 mV。闭环各项性能如表5-3所示。")
caption("表5-3  闭环特性测量结果")
table(
    ["项目", "测量条件", "测量值", "计算结果"],
    [
        ["电压放大倍数Avf", "Vi=46.7 mV，Vo=1.175 V", "—", "25.2"],
        ["反馈系数Fv", "Vf=36.4 mV，Vo=1.175 V", "—", "0.0310"],
        ["1/Fv", "—", "—", "32.3（理论26.5）"],
        ["输入电阻Ri", "R=200 kΩ，Vs=126.4 mV，Vi=104.5 mV", "—", "954 kΩ"],
        ["输出电阻Ro", "RL=200 Ω，Vo∞=1.43 V，VoL=1.36 V", "—", "10.3 Ω"],
        ["下限频率fL", "Vi=50 mV，Vom=1.012 V", "44 Hz", "—"],
        ["上限频率fH", "同上", "44.2 kHz", "—"],
        ["通频带fB", "—", "—", "44.16 kHz"],
    ],
    [3.0, 6.8, 2.6, 3.4], 9.5,
)
body("由表5-3可得：Avf=Vo/Vi=1.175/0.0467≈25.2；Fv=Vf/Vo=0.0364/1.175≈0.0310。若按反馈网络电阻计算，Fv≈RE1/(RE1+RF)=200/5300≈0.0377，1/Fv≈26.5。实测Avf=25.2与1/Fv=26.5的相对偏差约为4.9%，基本满足深度负反馈条件下Avf≈1/Fv的近似关系。")
body("实测Fv（0.0310）小于电阻分压理论值（0.0377），主要原因是理论公式Fv≈RE1/(RE1+RF)忽略了T1发射极看进去的晶体管等效电阻对RE1的并联作用，反馈网络实际上处于有载状态；此外RE1、RF的实际阻值存在±5%的容差，示波器读数也有一定误差。")
body("本实验开环增益A≈404，反馈系数F≈0.0377，环路增益AF≈15.2，反馈深度1+AF≈16.2。由此预测的闭环增益为A/(1+AF)≈24.9，与实测25.2仅相差约1.2%，进一步说明闭环增益主要由反馈网络决定。")
body("输入电阻：Ri=Vi/(Vs-Vi)×R=104.5/(126.4-104.5)×200 kΩ≈954 kΩ。与开环输入电阻50.7 kΩ相比增大了约18.8倍，方向与串联负反馈使输入电阻增大(1+AF)倍的规律一致；两者相差约16%，主要来自探头负载和读数误差。")
body("输出电阻：Ro=(Vo∞/VoL-1)×RL=(1.43/1.36-1)×200≈10.3 Ω。闭环输出电阻很小，接入200 Ω负载后输出电压仅下降约5%，因此示波器读数误差对结果影响较大，该结果宜作数量级参考。")

heading("5.3  开环特性", 2)
body("断开RF后，放大器处于开环状态。为避免高增益下输出波形失真，取Vi≈10 mV。测得开环特性如表5-4所示。")
caption("表5-4  开环特性测量结果")
table(
    ["项目", "测量条件", "测量值", "计算结果"],
    [
        ["电压放大倍数Av", "Vi=10.2 mV，Vo=4.12 V", "—", "404"],
        ["输入电阻Ri", "R=5.1 kΩ，Vs=100.6 mV，Vi=91.4 mV", "—", "50.7 kΩ"],
        ["输出电阻Ro", "RL=5.1 kΩ，Vo∞=5.72 V，VoL=4.12 V", "—", "1.98 kΩ"],
        ["下限频率fL", "Vi≈10 mV，Vom=4.12 V", "96 Hz", "—"],
        ["上限频率fH", "同上", "2.58 kHz", "—"],
        ["通频带fB", "—", "—", "2.48 kHz"],
    ],
    [3.0, 6.8, 2.6, 3.4], 9.5,
)
body("开环时Av=Vo/Vi=4.12/0.0102≈404；Ri=Vi/(Vs-Vi)×R=91.4/(100.6-91.4)×5.1 kΩ≈50.7 kΩ；Ro=(5.72/4.12-1)×5.1 kΩ≈1.98 kΩ。由于开环增益高，Vi仅10 mV左右时输出已达4.12 V，继续增大输入会使输出波形产生削顶失真，因此开环测量必须采用小信号。")

heading("5.4  开环与闭环性能对比", 2)
body("将开环与闭环的主要性能指标汇总于表5-5。")
caption("表5-5  开环与闭环性能对比")
table(
    ["性能指标", "开环", "闭环", "变化情况"],
    [
        ["电压放大倍数Av", "404", "25.2", "减小为约1/16"],
        ["输入电阻Ri", "50.7 kΩ", "954 kΩ", "增大18.8倍"],
        ["输出电阻Ro", "1.98 kΩ", "10.3 Ω", "显著减小"],
        ["下限频率fL", "96 Hz", "44 Hz", "降低"],
        ["上限频率fH", "2.58 kHz", "44.2 kHz", "提高17.1倍"],
        ["通频带fB", "2.48 kHz", "44.16 kHz", "扩大17.8倍"],
        ["增益带宽积GBW", "1.04 MHz", "1.11 MHz", "基本不变"],
    ],
    [3.8, 3.4, 3.4, 5.2],
)
body("由表5-5可知，引入电压串联负反馈后，电压放大倍数由404降至25.2，约为原来的1/16；输入电阻由50.7 kΩ增大到954 kΩ；输出电阻由1.98 kΩ减小到约10 Ω；通频带由2.48 kHz扩展到44.16 kHz，约扩大17.8倍。开环与闭环的增益带宽积分别为1.04 MHz和1.11 MHz，基本保持不变，符合增益带宽积为常数的规律。上述结果说明负反馈以降低增益为代价，换取了增益稳定性、输入电阻、输出电阻和通频带等性能的改善。")

heading("5.5  幅频特性曲线", 2)
body("根据表5-3、表5-4的中频增益和上、下限频率，绘出开环与闭环幅频特性曲线如图5-1所示。")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.add_run().add_picture(os.path.join(ASSETS, "bode.png"), width=Cm(15.0))
caption("图5-1  开环与闭环幅频特性曲线")
body("由图5-1可见，开环时中频增益约为52 dB，通频带较窄；闭环后中频增益降至约28 dB，但通频带明显展宽，下限频率由96 Hz降至44 Hz，上限频率由2.58 kHz提高到44.2 kHz。两条曲线的增益带宽积基本相等，直观地反映了负反馈“降增益、宽频带”的作用。")

heading("5.6  误差分析", 2)
errs = [
    "（1）万用表直流电压档存在±(0.5%+2字)的测量误差，且各电极电压是分别对地测量，计算VBE、VCE时误差会叠加；电阻的实际阻值与标称值存在±5%的容差，使由电压反算的电流和β产生偏差。",
    "（2）示波器测量的是峰峰值，mV级小信号的读数和光标定位误差较大；信号源的实际输出幅度随频率和负载变化，若未随时监视Vi，会使增益和通频带的测量产生误差。",
    "（3）示波器探头输入电阻约10 MΩ、输入电容约20 pF。对于200 Ω的RE1，10 MΩ输入电阻的分流可以忽略，因此Vf的偏差主要来自反馈网络的有载效应和读数误差；但探头电容在高阻节点会与节点电阻形成低通，对上限频率fH的测量有一定影响。",
    "（4）测量输入电阻时，串联电阻R与Ri处于同一数量级，Vi与Vs相差不大，两个相近电压相减会放大读数误差；示波器两通道探头的地线共地也可能引入额外分流。",
    "（5）测量输出电阻时，闭环Ro很小，接入200 Ω负载后输出电压仅下降约5%，电压变化量接近示波器的分辨极限，因此闭环Ro的相对误差较大；开环Ro较大，负载引起的电压变化明显，测量结果相对可靠。",
    "（6）频率测量受信号源频率显示精度、示波器时基误差以及电路分布电容、耦合电容的影响；在高、低频端波形可能失真，0.707Vom的判定存在一定主观误差。",
    "（7）T2的集电极电流由RC2支路算得约2.17 mA，而由RE2支路算得IE2≈2.28 mA，两者相差约5%，主要原因是RC2、RE2的阻值容差以及万用表读数误差。",
]
for s in errs:
    body(s, indent=False, space_after=2)

heading("5.7  实验结论", 2)
concl = [
    "（1）静态工作点测量结果表明，T1、T2的VBE均约为0.65 V、VCE均大于5 V，两管均工作在放大区，电路静态工作点基本正常。",
    "（2）闭环电压放大倍数实测为25.2，与由反馈网络理论值1/Fv≈26.5相差约4.9%，与由环路增益预测的24.9相差约1.2%，验证了深度负反馈条件下Avf≈1/Fv的近似关系，说明闭环增益主要由反馈网络决定，稳定性好。",
    "（3）电压串联负反馈使输入电阻由50.7 kΩ增大到954 kΩ，输出电阻由1.98 kΩ减小到约10 Ω，通频带由2.48 kHz扩展到44.16 kHz；开、闭环增益带宽积基本不变（约1.1 MHz）。",
    "（4）负反馈以牺牲电压增益为代价，全面改善了放大器的增益稳定性、输入电阻、输出电阻和频率响应，是模拟电路中最常用的性能改善手段。",
]
for s in concl:
    body(s, indent=False, space_after=2)

# =====================================================================
# 六、实验思考题
# =====================================================================
heading("六、实验思考题")

qa = [
    ("1. 负反馈为什么能提高放大倍数的稳定性？",
     "闭环增益Af=A/(1+AF)。对A求导可得 dAf/Af = [1/(1+AF)]·(dA/A)，即开环增益的相对变化被缩小(1+AF)倍。本实验AF≈15.2，1+AF≈16.2，因此晶体管参数、电源电压等变化引起的增益变化可减小到约1/16，增益稳定性明显提高。"),
    ("2. 电压串联负反馈对输入电阻和输出电阻有何影响？",
     "串联负反馈使反馈电压与输入电压串联相减，输入回路等效电阻增大，Rif=Ri(1+AF)；电压负反馈取样于输出电压，能维持输出电压稳定，使输出电阻减小，Rof=Ro/(1+AF)。本实验中Ri由50.7 kΩ增至954 kΩ，Ro由1.98 kΩ降至约10 Ω，与理论趋势一致。"),
    ("3. 测量输入电阻时，为什么串联电阻R要与Ri处于同一数量级？",
     "由Ri=Vi/(Vs-Vi)×R可知，若R远小于Ri，则Vs-Vi很小，两个相近电压相减的相对误差很大；若R远大于Ri，则Vi很小，信噪比和示波器分辨率不足。取R与Ri同数量级，可使Vi约为Vs的0.5～0.9，测量精度较高。因此闭环Ri高时选200 kΩ，开环Ri低时选5.1 kΩ。"),
    ("4. 测量输出电阻时，为什么负载电阻要与Ro相当？",
     "由Ro=(Vo∞/VoL-1)RL可知，若RL远大于Ro，接入负载后VoL几乎不变，电压变化量太小，读数误差大；若RL远小于Ro，则VoL过小，同样不易准确测量。取RL与Ro同数量级，可使输出电压有明显且可测的变化。本实验闭环Ro小，用200 Ω；开环Ro大，用5.1 kΩ。"),
    ("5. 测量静态工作点时，为什么必须使输入信号为零？",
     "万用表DCV档测量的是直流平均值。若输入端存在交流信号，晶体管的非线性会对交流信号产生整流和平均效应，使直流读数偏离真实的静态值；输入信号过大时还会使工作点随信号摆动。因此测量前应将输入端对地短路，确保Vi=0。"),
    ("6. 深度负反馈条件Avf≈1/Fv的含义是什么？本实验是否满足？",
     "当环路增益AF>>1（工程上一般取AF≥10）时，1+AF≈AF，于是Avf=A/(1+AF)≈1/F，闭环增益仅由反馈网络决定，与开环增益及晶体管参数基本无关。本实验A≈404、F≈0.0377，AF≈15.2，1+AF≈16.2，由公式算得Avf≈24.9，实测25.2，与1/Fv≈26.5相差约5%，可以认为基本满足深度负反馈条件。"),
    ("7. 测量通频带时，为什么以0.707Vom为界，并要保持Vi不变？",
     "0.707Vom对应功率下降到中频值的一半，即-3 dB点，是通频带的定义。测量时必须保持输入电压幅值和波形不变，这样输出电压的下降才能完全归因于放大电路本身的频率响应，而不是信号源幅度的变化；实际测量中应随时用示波器监视Vi并微调信号源。"),
    ("8. 若实测闭环增益明显偏离1/Fv，可能有哪些原因？",
     "可能原因包括：开环增益不够大，AF较小，不满足深度负反馈条件；测量Vf时探头负载使Vf偏小，导致Fv测量值偏小；RE1、RF实际阻值存在容差；输入信号过大使输出波形失真；示波器读数误差；电路存在寄生反馈或自激振荡等。"),
]
for q, a in qa:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    r = p.add_run(q)
    set_font(r, "宋体", 12, True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.5
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.first_line_indent = Pt(24)
    p2.paragraph_format.keep_together = True
    r2 = p2.add_run("答：" + a)
    set_font(r2, "宋体", 12)

footer_page_number()
doc.save(OUT)
print("saved", OUT)
